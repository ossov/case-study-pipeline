import os
# import requests # No longer needed
import urllib.request # Added
import json
import csv
import re
from pathlib import Path  # Using pathlib for modern path management
import docx  # Required to read .docx files
from docx.shared import Pt, Inches # Added for formatting
from docx.enum.text import WD_ALIGN_PARAGRAPH # Added for formatting

# --- Configuration ---
# API and Model Settings
# Replaced old API vars with the Ollama config
OLLAMA_URL = "http://localhost:11434/api/generate"
OLLAMA_MODEL = "deepseek-v3.1:671b-cloud" # Using model from your example
OLLAMA_TIMEOUT = 180 # Using timeout from your example

# --- File Path Configuration ---
# Use Path() to handle file paths correctly
# Paths are taken from the environment so the script carries no local layout.
#   CASE_STUDY_INPUT_DIR   raw practitioner notes (.docx)
#   CASE_STUDY_OUTPUT_DIR  generated narrative case studies (.docx)
#   CASE_STUDY_BASE_DIR    where story_log.csv is written
BASE_DIR = Path(os.environ.get("CASE_STUDY_BASE_DIR", "."))
INPUT_DIR = Path(os.environ.get("CASE_STUDY_INPUT_DIR", BASE_DIR / "sample_input"))
STORIES_OUTPUT_DIR = Path(os.environ.get("CASE_STUDY_OUTPUT_DIR", BASE_DIR / "output"))

# The spreadsheet will be created in the BASE_DIR
SPREADSHEET_FILENAME = BASE_DIR / "story_log.csv"
# --- End Configuration ---


# --- The Core Prompt: Your Instructions for the AI ---
SYSTEM_PROMPT = """
You are an expert writer and data extractor for a flower essence practitioner. Your specialty is transforming clinical notes into relatable, professional case studies.
Your task is to transform raw case study notes into an engaging narrative AND
extract key metadata from the notes.

You MUST return your response as a single, valid JSON object.
Do not include any text outside of the JSON object.

The JSON object must have the following exact structure:

{
  "story_title": "A clear, professional, and SEO-friendly title. Focus *only* on the client's main symptom or core challenge (e.g., 'Chronic Fatigue', 'Navigating Deep-Seated Grief', 'Adolescent Acne'). The title should be short, direct, and just 2-5 words. Avoid 'From X to Y' patterns. Do NOT use the client's name.",
  "main_symptom": "A very brief (3-7 word) summary of the client's chief complaint (e.g., 'Feeling stuck and afraid of failure', 'Grief and emotional numbness').",
  "essences_used": [
    "A list of all unique essence names mentioned.",
    "List each essence only once, even if used in multiple sessions.",
    "Make sure to capture all of them."
  ],
  "story_body": "The complete, formatted story text. Follow all these rules:
    1.  **Tone:** Write in a professional, empathetic, and human-centered tone. The style should be grounded and credible, like a journalistic feature, focusing on the client's real-world struggles and the practical changes they experienced. Avoid overly dramatic or 'fairy tale' phrasing. This is a real person's journey.
    2.  **Format:** Start with the title as an '###' heading.
    3.  **Privacy Note:** On the next line, add the italicized note: *This story is based on a real client's experience. Names and identifying details have been changed to protect her privacy.*
    4.  **Narrative:** Write the story in the third person ('she felt...').
    5.  **Structure:** Follow the 5-part arc: The 'Before' state, The 'Turning Point', The 'Journey' (mentioning the essences), The 'Breakthrough', and The 'Transformed After'.
    6.  **Essences:** When you mention the essences in the story's text, make their names **bold**.
    7.  **Length:** The final story should be 400-800 words."
}
"""

def read_docx_text(filepath: Path) -> str | None:
    """Extracts all text from a .docx file."""
    try:
        doc = docx.Document(filepath)
        full_text = [para.text for para in doc.paragraphs]
        return '\n'.join(full_text)
    except Exception as e:
        print(f"    Error reading {filepath.name}: {e}")
        return None

def sanitize_filename(title: str) -> str:
    """Turns a story title into a safe filename (e.g., 'story-title.docx')."""
    s = title.lower()
    s = re.sub(r'[^\w\s-]', '', s)  # Remove non-alphanumeric chars
    s = re.sub(r'[\s_]+', '-', s)   # Replace spaces with hyphens
    s = s.strip('-')
    return f"{s}.docx" # Changed to .docx

def log_to_spreadsheet(original_filename, new_story_filename, main_symptom, essences_used):
    """
    Appends a new row of metadata to the master CSV file.
    Creates the file and header row if it doesn't exist.
    """
    file_exists = SPREADSHEET_FILENAME.is_file()
    
    # Join the list of essences into a single string for the CSV
    essences_str = ", ".join(sorted(essences_used))
    
    row_data = [original_filename, new_story_filename, main_symptom, essences_str]
    
    try:
        with open(SPREADSHEET_FILENAME, 'a', newline='', encoding='utf-8') as f:
            writer = csv.writer(f)
            
            if not file_exists:
                # Write the header row
                header = ["original_file", "new_story_file", "main_symptom", "essences_used"]
                writer.writerow(header)
            
            # Write the new data row
            writer.writerow(row_data)
            
        print(f"    Successfully logged metadata to {SPREADSHEET_FILENAME.name}")
        
    except IOError as e:
        print(f"    Error: Could not write to CSV file. {e}")

def generate_story_from_case_study(case_study_notes: str, original_filename: str) -> str:
    """
    Takes raw case study notes, generates a story, saves it,
    and logs the metadata using the Ollama API format.
    """
    # Removed the API_KEY check

    # Construct the prompt in the same way organize_case_studies.py does
    full_prompt = (
        f"{SYSTEM_PROMPT}\n\n"
        f"Here are the case study notes. Generate the JSON output:\n\n"
        f"---\n\n{case_study_notes}"
    )
    
    try:
        print("    Contacting AI... (This may take a moment)")

        # --- New API Call Logic (from organize_case_studies.py) ---
        payload = {
            'model': OLLAMA_MODEL,
            'prompt': full_prompt,
            'stream': False,
            'options': {'temperature': 0.7} # Kept original temp setting
        }
        
        data = json.dumps(payload).encode('utf-8')
        req = urllib.request.Request(
            OLLAMA_URL,
            data=data,
            headers={'Content-Type': 'application/json'}
        )
        
        with urllib.request.urlopen(req, timeout=OLLAMA_TIMEOUT) as response:
            result = json.loads(response.read().decode())
            response_text = result['response'].strip()
        # --- End of New API Call Logic ---

        try:
            # Handle cases where the model wraps the JSON in markdown
            if response_text.startswith("```json"):
                response_text = re.sub(r"^```json\n|```$", "", response_text, flags=re.MULTILINE)

            data = json.loads(response_text)
            
            story_title = data['story_title']
            story_body = data['story_body']
            main_symptom = data['main_symptom']
            essences_used = data['essences_used']

        except json.JSONDecodeError as e:
            return f"    Error: Failed to decode JSON from AI response. {e}\n    Response was:\n{response_text}"
        except KeyError as e:
            return f"    Error: Missing expected key {e} in AI's JSON response.\n    Response was:\n{response_text}"
        
        # --- File Saving (Changed to .docx) ---
        new_story_filename = sanitize_filename(story_title)
        # Save the new story to the specified output directory
        output_path = STORIES_OUTPUT_DIR / new_story_filename
        
        try:
            doc = docx.Document()
            # Set default font
            style = doc.styles['Normal']
            font = style.font
            font.name = 'Arial'
            font.size = Pt(11)

            lines = story_body.split('\n')
            
            for line in lines:
                
                if line.startswith('###'):
                    text = line.replace('###', '').strip()
                    heading = doc.add_heading(level=1)
                    run = heading.add_run(text)
                    run.font.name = 'Arial'
                    run.font.size = Pt(16)
                    run.bold = True
                
                elif line.startswith('*') and line.endswith('*') and not '**' in line: # Italic privacy note
                    text = line.strip('*').strip()
                    p = doc.add_paragraph()
                    p.add_run(text).italic = True
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER 
                
                # Removed the blockquote handler
                # elif line.startswith('>'): ...
                
                elif not line.strip(): # Empty line
                    doc.add_paragraph()
                
                elif line: # Normal paragraph with potential bolding
                    p = doc.add_paragraph()
                    # Split by bold markers, preserving them
                    parts = re.split(r'(\*\*.*?\*\*)', line)
                    for part in parts:
                        if part.startswith('**') and part.endswith('**'):
                            text = part.strip('*') # Get text inside **
                            p.add_run(text).bold = True
                        elif part: # Add normal text
                            p.add_run(part)

            doc.save(output_path)
            print(f"    Successfully saved story to {output_path}")
            
        except Exception as e:
            # Catch all errors during docx creation
            return f"    Error: Could not save .docx story file. {e}"

        # --- Metadata Logging (unchanged) ---
        log_to_spreadsheet(original_filename, new_story_filename, main_symptom, essences_used)

        return f"    Success! Story created: {new_story_filename}"

    # --- New Error Handling (from organize_case_studies.py) ---
    except urllib.error.URLError as e:
        if "timed out" in str(e): # Simple check for timeout
            print(f"    ❌ OLLAMA ERROR: The request timed out after {OLLAMA_TIMEOUT} seconds.")
        else:
            print(f"    ❌ OLLAMA ERROR: Could not connect to the Ollama server at {OLLAMA_URL}.")
            print("       Please ensure the Ollama application is running.")
        return f"    Error making API call: {e}"
    except Exception as e:
        # Catch other potential errors (e.g., in json.loads if response is not JSON)
        return f"    An unexpected error occurred: {e}"


# --- Main Batch Processing Logic ---
if __name__ == "__main__":
    print("--- Starting Story Generation Batch ---")

    # 1. Removed the API Key check
    
    # 2. Create output directory if it doesn't exist
    try:
        STORIES_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        print(f"Stories will be saved to: {STORIES_OUTPUT_DIR}")
    except Exception as e:
        print(f"Error creating output directory {STORIES_OUTPUT_DIR}: {e}")
        exit()

    # 3. Find case study files
    print(f"Scanning for case studies in: {INPUT_DIR}")
    try:
        case_study_files = list(INPUT_DIR.glob("*.docx"))
    except Exception as e:
        print(f"Error scanning directory {INPUT_DIR}: {e}")
        exit()

    if not case_study_files:
        print(f"Error: No .docx files found in {INPUT_DIR}. Exiting.")
        exit()

    # 4. Process all files
    files_to_process = case_study_files
    print(f"Found {len(case_study_files)} total files. Processing all of them.")

    # 5. Loop through and process each file
    for i, docx_path in enumerate(files_to_process, 1):
        original_file = docx_path.name
        print(f"\n--- Processing file {i}/{len(files_to_process)}: {original_file} ---")
        
        case_study_notes = read_docx_text(docx_path)
        
        if case_study_notes:
            # If text was read successfully, generate the story
            result_message = generate_story_from_case_study(
                case_study_notes=case_study_notes,
                original_filename=original_file
            )
            print(result_message)
        else:
            # If read_docx_text returned None
            print(f"    Skipping {original_file} due to read error.")
            
    print("\n--- Batch Processing Complete ---")
    print(f"Check {STORIES_OUTPUT_DIR} for your stories.")
    print(f"Check {SPREADSHEET_FILENAME} for the metadata log.")